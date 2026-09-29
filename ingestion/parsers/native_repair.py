"""使用源文件的机器可读结构修复确定性解析缺口。"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from ingestion.document_model import (
    DocumentElement,
    DocumentFormat,
    ElementType,
    IssueSeverity,
    ParseIssue,
    ParsedDocument,
    SourceLocation,
    TableCell,
    TableData,
    stable_element_id,
)
from ingestion.parsers.review_models import ReviewTarget


_COMPACT_RE = re.compile(r"[^0-9A-Za-z\u3400-\u9fff]+")


class NativeStructureRepairer:
    """只处理不需要视觉判断的目标，并为每个补回元素保留来源定位。"""

    def repair(
        self,
        document: ParsedDocument,
        source_path: Path,
        targets: list[ReviewTarget],
    ) -> tuple[ParsedDocument, list[dict[str, Any]]]:
        deterministic = [target for target in targets if not target.requires_vision]
        if not deterministic:
            return document, []

        elements = list(document.elements)
        findings: list[dict[str, Any]] = []
        for target in deterministic:
            before = len(elements)
            if document.source.document_format == DocumentFormat.PDF:
                self._repair_pdf(elements, document, source_path, target)
            elif document.source.document_format == DocumentFormat.DOCX:
                self._repair_docx(elements, document, source_path, target)
            elif document.source.document_format == DocumentFormat.PPTX:
                self._repair_pptx(elements, document, source_path, target)
            elif document.source.document_format == DocumentFormat.XLSX:
                self._repair_xlsx(elements, document, source_path, target)
            added = len(elements) - before
            findings.append({
                "target_id": target.target_id,
                "locator": target.locator,
                "status": "repaired" if added else "no_change",
                "added_element_count": added,
                "method": f"native_{document.source.document_format.value}",
            })

        added_total = len(elements) - len(document.elements)
        issue = ParseIssue(
            code="NATIVE_STRUCTURE_REPAIR",
            message=f"原生结构修复器为 {len(deterministic)} 个目标补回 {added_total} 个元素",
            severity=IssueSeverity.INFO,
            retryable=False,
            details={"findings": findings},
        )
        issues = [value for value in document.issues if value.code != issue.code] + [issue]
        repaired = document.model_copy(update={
            "elements": elements,
            "root_element_ids": [element.element_id for element in elements],
            "issues": issues,
        })
        return repaired, findings

    def _repair_pdf(
        self, elements: list[DocumentElement], document: ParsedDocument,
        source_path: Path, target: ReviewTarget,
    ) -> None:
        if not target.page_number:
            return
        existing = _unit_text(elements, page_number=target.page_number)
        from pypdf import PdfReader

        reader = PdfReader(str(source_path))
        if not 1 <= target.page_number <= len(reader.pages):
            return
        text = _clean_text(reader.pages[target.page_number - 1].extract_text() or "")
        if not _should_add(text, existing):
            return
        locator = f"native-repair:page:{target.page_number}:full-text"
        elements.append(_element(
            document, elements, ElementType.PARAGRAPH, text, locator,
            SourceLocation(page_number=target.page_number, parser_locator=locator), target,
            metadata={"layout_mode": "native_full_text_fallback"},
        ))

    def _repair_docx(
        self, elements: list[DocumentElement], document: ParsedDocument,
        source_path: Path, target: ReviewTarget,
    ) -> None:
        from docx import Document

        source = Document(str(source_path))
        existing = _unit_text(elements)
        for index, paragraph in enumerate(source.paragraphs):
            text = _clean_text(paragraph.text)
            if not _should_add(text, existing):
                continue
            locator = f"native-repair:paragraph:{index}"
            elements.append(_element(
                document, elements, ElementType.PARAGRAPH, text, locator,
                SourceLocation(paragraph_index=index, parser_locator=locator), target,
            ))
            existing += "\n" + text
        for table_index, table in enumerate(source.tables):
            for row_index, row in enumerate(table.rows):
                text = "\t".join(_clean_text(cell.text) for cell in row.cells if _clean_text(cell.text))
                if not _should_add(text, existing):
                    continue
                locator = f"native-repair:table:{table_index}:row:{row_index}"
                elements.append(_element(
                    document, elements, ElementType.PARAGRAPH, text, locator,
                    SourceLocation(
                        table_index=table_index, row_index=row_index, parser_locator=locator,
                    ), target,
                ))
                existing += "\n" + text

    def _repair_pptx(
        self, elements: list[DocumentElement], document: ParsedDocument,
        source_path: Path, target: ReviewTarget,
    ) -> None:
        from pptx import Presentation

        if not target.slide_number:
            return
        presentation = Presentation(str(source_path))
        if not 1 <= target.slide_number <= len(presentation.slides):
            return
        existing = _unit_text(elements, slide_number=target.slide_number)
        slide = presentation.slides[target.slide_number - 1]
        for shape_index, shape in enumerate(slide.shapes):
            candidates: list[tuple[str, str]] = []
            if getattr(shape, "has_text_frame", False):
                candidates.append(("text", _clean_text(shape.text)))
            if getattr(shape, "has_table", False):
                for row_index, row in enumerate(shape.table.rows):
                    candidates.append((
                        f"table-row:{row_index}",
                        "\t".join(_clean_text(cell.text) for cell in row.cells if _clean_text(cell.text)),
                    ))
            for suffix, text in candidates:
                if not _should_add(text, existing):
                    continue
                locator = f"native-repair:slide:{target.slide_number}:shape:{shape_index}:{suffix}"
                elements.append(_element(
                    document, elements, ElementType.PARAGRAPH, text, locator,
                    SourceLocation(slide_number=target.slide_number, parser_locator=locator), target,
                ))
                existing += "\n" + text

    def _repair_xlsx(
        self, elements: list[DocumentElement], document: ParsedDocument,
        source_path: Path, target: ReviewTarget,
    ) -> None:
        from openpyxl import load_workbook

        if not target.sheet_name:
            return
        workbook = load_workbook(source_path, read_only=False, data_only=False)
        values = load_workbook(source_path, read_only=False, data_only=True)
        try:
            if target.sheet_name not in workbook.sheetnames:
                return
            sheet = workbook[target.sheet_name]
            value_sheet = values[target.sheet_name]
            cells: list[TableCell] = []
            max_row = 0
            max_column = 0
            for row in sheet.iter_rows():
                for cell in row:
                    displayed = value_sheet[cell.coordinate].value
                    formula = cell.value if isinstance(cell.value, str) and cell.value.startswith("=") else None
                    raw_value = cell.value
                    if displayed is None and formula is None:
                        displayed = raw_value
                    if displayed is None and formula is None:
                        continue
                    max_row = max(max_row, cell.row)
                    max_column = max(max_column, cell.column)
                    cells.append(TableCell(
                        row=cell.row - 1,
                        column=cell.column - 1,
                        text="" if displayed is None else str(displayed),
                        raw_value=raw_value if isinstance(raw_value, (str, int, float, bool)) else None,
                        formula=formula,
                        is_header=cell.row == 1,
                    ))
            if not cells:
                return
            locator = f"native-repair:sheet:{sheet.title}:used-range"
            order = len(elements)
            elements.append(DocumentElement(
                element_id=stable_element_id(document.source.document_id, ElementType.TABLE, order, locator),
                element_type=ElementType.TABLE,
                order=order,
                text=sheet.title,
                source_location=SourceLocation(
                    sheet_name=sheet.title,
                    cell_range=f"A1:{sheet.cell(max_row, max_column).coordinate}",
                    parser_locator=locator,
                ),
                table=TableData(row_count=max_row, column_count=max_column, cells=cells, caption=sheet.title),
                metadata={
                    "source": "native_structure_repair",
                    "target_id": target.target_id,
                    "repair_anomalies": [value.value for value in target.anomalies],
                    "formula_count": sum(cell.formula is not None for cell in cells),
                    "layout_mode": "native_worksheet_used_range",
                },
            ))
        finally:
            workbook.close()
            values.close()


def _element(
    document: ParsedDocument,
    elements: list[DocumentElement],
    kind: ElementType,
    text: str,
    locator: str,
    location: SourceLocation,
    target: ReviewTarget,
    metadata: dict[str, Any] | None = None,
) -> DocumentElement:
    order = len(elements)
    return DocumentElement(
        element_id=stable_element_id(document.source.document_id, kind, order, locator),
        element_type=kind,
        order=order,
        text=text,
        source_location=location,
        metadata={
            "source": "native_structure_repair",
            "target_id": target.target_id,
            "repair_anomalies": [value.value for value in target.anomalies],
            **(metadata or {}),
        },
    )


def _unit_text(
    elements: list[DocumentElement],
    page_number: int | None = None,
    slide_number: int | None = None,
) -> str:
    selected: list[DocumentElement] = []
    for element in elements:
        location = element.source_location
        if page_number is not None and (not location or location.page_number != page_number):
            continue
        if slide_number is not None and (not location or location.slide_number != slide_number):
            continue
        selected.append(element)
    return "\n".join(
        [element.text for element in selected]
        + [cell.text for element in selected if element.table for cell in element.table.cells]
    )


def _should_add(candidate: str, existing: str) -> bool:
    compact = _compact(candidate)
    if len(compact) < 2:
        return False
    return compact not in _compact(existing)


def _compact(value: str) -> str:
    return _COMPACT_RE.sub("", value).lower()


def _clean_text(value: str) -> str:
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in value.replace("\r", "\n").split("\n")]
    return "\n".join(line for line in lines if line)


__all__ = ["NativeStructureRepairer"]
