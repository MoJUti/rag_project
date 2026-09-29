"""使用文件原生结构生成独立的完整性参考。"""

from __future__ import annotations

import re
from html.parser import HTMLParser
from pathlib import Path

from pydantic import BaseModel, ConfigDict, Field

from ingestion.document_model import DocumentFormat


_NUMBER_RE = re.compile(
    r"(?<![0-9A-Za-z_.])[-+]?\d+(?:(?:,|，)\s*\d{3})*(?:\.\d+)?%?(?![0-9A-Za-z_.])"
)


class NativeUnit(BaseModel):
    model_config = ConfigDict(extra="forbid")

    locator: str
    text: str = ""
    numeric_tokens: set[str] = Field(default_factory=set)
    table_count: int = 0
    formula_count: int = 0
    image_count: int = 0
    chart_count: int = 0


class NativeDocumentSnapshot(BaseModel):
    model_config = ConfigDict(extra="forbid")

    document_format: DocumentFormat
    units: list[NativeUnit] = Field(default_factory=list)
    table_count: int = 0
    formula_count: int = 0
    merged_range_count: int = 0
    metadata: dict[str, int | str | list[str]] = Field(default_factory=dict)

    @property
    def text(self) -> str:
        return "\n".join(unit.text for unit in self.units)

    @property
    def numeric_tokens(self) -> set[str]:
        result: set[str] = set()
        for unit in self.units:
            result.update(unit.numeric_tokens)
        return result


def extract_native_snapshot(path: Path, document_format: DocumentFormat) -> NativeDocumentSnapshot:
    if document_format == DocumentFormat.PDF:
        return _extract_pdf(path)
    if document_format == DocumentFormat.DOCX:
        return _extract_docx(path)
    if document_format == DocumentFormat.PPTX:
        return _extract_pptx(path)
    if document_format == DocumentFormat.XLSX:
        return _extract_xlsx(path)
    if document_format == DocumentFormat.HTML:
        return _extract_html(path)
    return NativeDocumentSnapshot(document_format=document_format)


def numeric_tokens(text: str) -> set[str]:
    return {_normalize_number(match.group(0)) for match in _NUMBER_RE.finditer(text)}


def missing_numeric_tokens(expected: set[str], parsed_text: str) -> set[str]:
    """允许千分位和数字内部空格差异，同时避免子串误匹配。"""
    parsed_tokens = numeric_tokens(parsed_text)
    compact = parsed_text.replace(",", "").replace("，", "")
    compact = re.sub(r"(?<=\d)\s+(?=\.)", "", compact)
    compact = re.sub(r"(?<=\.)\s+(?=\d)", "", compact)
    missing: set[str] = set()
    for token in expected:
        if token in parsed_tokens:
            continue
        pattern = rf"(?<![\d.]){re.escape(token)}(?![\d.])"
        if re.search(pattern, compact) is None:
            missing.add(token)
    return missing


def _normalize_number(value: str) -> str:
    normalized = re.sub(r"[,，\s]", "", value).lstrip("+")
    if normalized.endswith("%"):
        normalized = normalized[:-1]
    if "." in normalized:
        normalized = normalized.rstrip("0").rstrip(".")
    return normalized


def _unit(locator: str, text: str, **counts: int) -> NativeUnit:
    return NativeUnit(locator=locator, text=text, numeric_tokens=numeric_tokens(text), **counts)


def _extract_pdf(path: Path) -> NativeDocumentSnapshot:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    units = [_unit(f"page:{index + 1}", page.extract_text() or "") for index, page in enumerate(reader.pages)]
    return NativeDocumentSnapshot(
        document_format=DocumentFormat.PDF,
        units=units,
        metadata={"page_count": len(reader.pages)},
    )


def _extract_docx(path: Path) -> NativeDocumentSnapshot:
    from docx import Document

    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        parts.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
    text = "\n".join(parts)
    return NativeDocumentSnapshot(
        document_format=DocumentFormat.DOCX,
        units=[_unit("document", text, table_count=len(document.tables))],
        table_count=len(document.tables),
        metadata={"paragraph_count": len(document.paragraphs)},
    )


def _extract_pptx(path: Path) -> NativeDocumentSnapshot:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("解析 PPTX 需要安装 python-pptx") from exc

    presentation = Presentation(str(path))
    units: list[NativeUnit] = []
    table_count = 0
    for slide_index, slide in enumerate(presentation.slides, start=1):
        parts: list[str] = []
        slide_tables = 0
        slide_images = 0
        slide_charts = 0
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                parts.append(shape.text)
            if getattr(shape, "has_table", False):
                table_count += 1
                slide_tables += 1
                parts.extend("\t".join(cell.text for cell in row.cells) for row in shape.table.rows)
            if getattr(shape, "has_chart", False):
                slide_charts += 1
            if getattr(shape, "shape_type", None) == 13:  # MSO_SHAPE_TYPE.PICTURE
                slide_images += 1
        units.append(_unit(
            f"slide:{slide_index}", "\n".join(parts), table_count=slide_tables,
            image_count=slide_images, chart_count=slide_charts,
        ))
    return NativeDocumentSnapshot(
        document_format=DocumentFormat.PPTX,
        units=units,
        table_count=table_count,
        metadata={"slide_count": len(presentation.slides)},
    )


def _extract_xlsx(path: Path) -> NativeDocumentSnapshot:
    from openpyxl import load_workbook

    formulas = load_workbook(path, read_only=False, data_only=False)
    values = load_workbook(path, read_only=False, data_only=True)
    units: list[NativeUnit] = []
    formula_count = 0
    merged_count = 0
    try:
        for formula_sheet in formulas.worksheets:
            value_sheet = values[formula_sheet.title]
            parts: list[str] = []
            sheet_formulas = 0
            merged_count += len(formula_sheet.merged_cells.ranges)
            for row in formula_sheet.iter_rows():
                row_values: list[str] = []
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formula_count += 1
                        sheet_formulas += 1
                    displayed = value_sheet[cell.coordinate].value
                    if displayed is None and not (isinstance(cell.value, str) and cell.value.startswith("=")):
                        displayed = cell.value
                    if displayed is not None:
                        row_values.append(str(displayed))
                if row_values:
                    parts.append("\t".join(row_values))
            units.append(_unit(
                f"sheet:{formula_sheet.title}", "\n".join(parts),
                formula_count=sheet_formulas, table_count=len(formula_sheet.tables),
            ))
    finally:
        formulas.close()
        values.close()
    return NativeDocumentSnapshot(
        document_format=DocumentFormat.XLSX,
        units=units,
        formula_count=formula_count,
        merged_range_count=merged_count,
        metadata={"sheet_count": len(units), "sheet_names": [unit.locator[6:] for unit in units]},
    )


class _NativeHtmlParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.table_count = 0
        self._ignored_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style", "noscript"}:
            self._ignored_depth += 1
        if tag == "table":
            self.table_count += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript"} and self._ignored_depth:
            self._ignored_depth -= 1

    def handle_data(self, data: str) -> None:
        if not self._ignored_depth and data.strip():
            self.parts.append(data.strip())


def _extract_html(path: Path) -> NativeDocumentSnapshot:
    parser = _NativeHtmlParser()
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    text = "\n".join(parser.parts)
    return NativeDocumentSnapshot(
        document_format=DocumentFormat.HTML,
        units=[_unit("document", text, table_count=parser.table_count)],
        table_count=parser.table_count,
    )


__all__ = [
    "NativeDocumentSnapshot",
    "NativeUnit",
    "extract_native_snapshot",
    "missing_numeric_tokens",
    "numeric_tokens",
]
