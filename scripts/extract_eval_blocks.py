"""Extract location-aware text blocks from the curated evaluation documents."""

from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import fitz
import openpyxl
from docx import Document


ROOT = Path(__file__).resolve().parents[1]
BUNDLE = ROOT / "data" / "eval" / "zh_public_minibench_v0.1"
MANIFEST = BUNDLE / "manifest.jsonl"
OUTPUT = BUNDLE / "annotations" / "extracted_blocks.jsonl"


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def docx_blocks(path: Path) -> list[dict]:
    document = Document(path)
    blocks = []
    for index, paragraph in enumerate(document.paragraphs, 1):
        text = clean(paragraph.text)
        if text:
            blocks.append({"locator": {"type": "paragraph", "index": index}, "text": text})
    for table_index, table in enumerate(document.tables, 1):
        for row_index, row in enumerate(table.rows, 1):
            values = [clean(cell.text) for cell in row.cells]
            text = " | ".join(value for value in values if value)
            if text:
                blocks.append(
                    {
                        "locator": {
                            "type": "table_row",
                            "table": table_index,
                            "row": row_index,
                        },
                        "text": text,
                    }
                )
    if blocks:
        return blocks

    # Some benchmark DOCX files place every visible string in drawing/text-box
    # XML, which python-docx does not expose through document.paragraphs.
    word_namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
        for index, paragraph in enumerate(root.iter(word_namespace + "p"), 1):
            text = clean(
                " ".join(node.text or "" for node in paragraph.iter(word_namespace + "t"))
            )
            if text:
                blocks.append(
                    {"locator": {"type": "xml_paragraph", "index": index}, "text": text}
                )
    return blocks


def pptx_blocks(path: Path) -> list[dict]:
    namespace = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    blocks = []
    with zipfile.ZipFile(path) as archive:
        slides = sorted(
            (
                name
                for name in archive.namelist()
                if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
            ),
            key=lambda name: int(re.search(r"(\d+)", Path(name).stem).group(1)),
        )
        for slide_number, name in enumerate(slides, 1):
            root = ElementTree.fromstring(archive.read(name))
            text = clean(" ".join(node.text or "" for node in root.iter(namespace + "t")))
            if text:
                blocks.append(
                    {"locator": {"type": "slide", "slide": slide_number}, "text": text}
                )
    return blocks


def xlsx_blocks(path: Path) -> list[dict]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    blocks = []
    for worksheet in workbook.worksheets:
        for row_number, row in enumerate(worksheet.iter_rows(), 1):
            values = []
            cells = []
            for cell in row:
                value = clean(cell.value)
                if value:
                    values.append(value)
                    cells.append(cell.coordinate)
            if values:
                blocks.append(
                    {
                        "locator": {
                            "type": "sheet_row",
                            "sheet": worksheet.title,
                            "row": row_number,
                            "cells": cells,
                        },
                        "text": " | ".join(values),
                    }
                )
    workbook.close()
    return blocks


def pdf_blocks(path: Path) -> list[dict]:
    document = fitz.open(path)
    return [
        {
            "locator": {"type": "page", "page": page.number + 1},
            "text": clean(page.get_text("text")),
        }
        for page in document
        if clean(page.get_text("text"))
    ]


def markdown_blocks(path: Path) -> list[dict]:
    return [
        {"locator": {"type": "line", "line": index}, "text": clean(line)}
        for index, line in enumerate(path.read_text(encoding="utf-8").splitlines(), 1)
        if clean(line)
    ]


def main() -> None:
    manifest = [json.loads(line) for line in MANIFEST.read_text(encoding="utf-8").splitlines()]
    records = []
    for item in manifest:
        path = BUNDLE / item["path"]
        extension = path.suffix.lower()
        if extension == ".docx":
            blocks = docx_blocks(path)
        elif extension == ".pptx":
            blocks = pptx_blocks(path)
        elif extension == ".xlsx":
            blocks = xlsx_blocks(path)
        elif extension == ".pdf":
            blocks = pdf_blocks(path)
        elif extension == ".md":
            blocks = markdown_blocks(path)
        else:
            continue
        for block_index, block in enumerate(blocks, 1):
            records.append(
                {
                    "document_id": item["item_id"],
                    "block_id": f"{item['item_id']}#b{block_index:05d}",
                    **block,
                }
            )
    with OUTPUT.open("w", encoding="utf-8", newline="\n") as stream:
        for record in records:
            stream.write(json.dumps(record, ensure_ascii=False) + "\n")
    print(json.dumps({"documents": len({r['document_id'] for r in records}), "blocks": len(records)}))


if __name__ == "__main__":
    main()
